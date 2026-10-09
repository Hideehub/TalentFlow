import re
from collections import Counter
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.db.models.functions import Lower
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from accounts.tenancy import get_user_company
from recruitment.choices import APPLICATION_STATUS, CANDIDATE_SOURCE
from recruitment.models import Application, ApplicationStatusChange, Candidate, JobOpening

from .choices import (
    CANDIDATE_HEADER_ALIASES,
    MAX_IMPORT_ROWS,
    REQUIRED_CANDIDATE_FIELDS,
)


class CandidateImportError(Exception):
    pass


def _normalize_header(value):
    text = str(value or "").strip().lower()
    return re.sub(r"\s+", " ", text)


def _header_map(header_row):
    normalized_aliases = {
        _normalize_header(alias): field
        for field, aliases in CANDIDATE_HEADER_ALIASES.items()
        for alias in aliases
    }

    return {
        normalized_aliases[_normalize_header(value)]: index
        for index, value in enumerate(header_row)
        if _normalize_header(value) in normalized_aliases
    }


def _select_worksheet(workbook):
    for worksheet in workbook.worksheets:
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        mapping = _header_map(header_row)
        if all(field in mapping for field in REQUIRED_CANDIDATE_FIELDS):
            return worksheet, mapping

    raise CandidateImportError("No sheet contains all required candidate columns.")


def _cell_value(row, mapping, field):
    index = mapping.get(field)
    if index is None:
        return None
    return row[index] if index < len(row) else None


def _normalize_name(value):
    name = str(value or "").strip()
    return name.title() if name.isupper() or name.islower() else name


def _normalize_status(value):
    return str(value or "").strip().lower().replace(" ", "_")


def _normalize_source(value):
    """Match a source by code or label; blank means "Excel import", unknown gives None."""
    text = str(value or "").strip().lower()
    if not text:
        return "import"
    for code, label in CANDIDATE_SOURCE:
        if text in (code, label.lower(), code.replace("_", " ")):
            return code
    return None


def _normalize_experience(value):
    if value in (None, ""):
        return None

    try:
        number = int(float(value))
    except (TypeError, ValueError):
        return None

    return number if number >= 0 else None


def _candidates_by_email(company, emails):
    return {
        candidate.email.lower(): candidate
        for candidate in Candidate.objects.annotate(email_key=Lower("email")).filter(
            company=company, email_key__in=emails
        )
    }


def _open_jobs_by_title(company):
    return {
        job.title.lower(): job
        for job in JobOpening.objects.filter(status="open", company=company)
    }


def _application_key(email, job_id, position):
    # A matched job identifies the application; without one, the imported position text does.
    return (email, ("job", job_id)) if job_id else (email, ("position", position.lower()))


def _existing_application_keys(company, emails):
    rows = (
        Application.objects.annotate(
            email_key=Lower("candidate__email"),
            position_key=Lower("imported_position"),
        )
        .filter(candidate__company=company, email_key__in=emails)
        .values_list("email_key", "job_id", "position_key")
    )
    return {_application_key(email, job_id, position) for email, job_id, position in rows}


def candidate_import_preview(uploaded_file, *, user=None):
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise CandidateImportError("The uploaded Excel file could not be read.") from exc

    worksheet, mapping = _select_worksheet(workbook)
    raw_rows = list(
        worksheet.iter_rows(
            min_row=2,
            max_row=MAX_IMPORT_ROWS + 2,
            values_only=True,
        )
    )
    non_empty_rows = [row for row in raw_rows if any(value not in (None, "") for value in row)]
    workbook.close()

    if len(non_empty_rows) > MAX_IMPORT_ROWS:
        raise CandidateImportError(f"Imports are limited to {MAX_IMPORT_ROWS} rows at a time.")

    parsed_rows = []
    for row_number, row in enumerate(non_empty_rows, start=2):
        email = str(_cell_value(row, mapping, "email") or "").strip().lower()
        parsed_rows.append(
            {
                "row_number": row_number,
                "full_name": _normalize_name(_cell_value(row, mapping, "full_name")),
                "email": email,
                "phone": str(_cell_value(row, mapping, "phone") or "").strip(),
                "position": str(_cell_value(row, mapping, "position") or "").strip(),
                "years_of_experience": _normalize_experience(
                    _cell_value(row, mapping, "years_of_experience")
                ),
                "status": _normalize_status(_cell_value(row, mapping, "status")),
                "source": _normalize_source(_cell_value(row, mapping, "source")),
                "errors": [],
            }
        )

    # The same person may appear on several rows, but only once per position.
    email_position_counts = Counter(
        (row["email"], row["position"].lower()) for row in parsed_rows if row["email"]
    )
    company = get_user_company(user)
    emails = [row["email"] for row in parsed_rows if row["email"]]
    existing_candidates = _candidates_by_email(company, emails)
    existing_keys = _existing_application_keys(company, emails)
    jobs_by_title = _open_jobs_by_title(company)
    valid_statuses = {value for value, _label in APPLICATION_STATUS}
    source_labels = dict(CANDIDATE_SOURCE)

    for row in parsed_rows:
        if not row["full_name"]:
            row["errors"].append("Full name is required.")

        if not row["email"]:
            row["errors"].append("Email is required.")
        else:
            try:
                validate_email(row["email"])
            except ValidationError:
                row["errors"].append("Email is invalid.")

            if email_position_counts[(row["email"], row["position"].lower())] > 1:
                row["errors"].append("Same email and position appear more than once in this file.")

        job = jobs_by_title.get(row["position"].lower())
        row["job_title"] = job.title if job else ""
        row["existing_candidate"] = row["email"] in existing_candidates
        if _application_key(row["email"], job and job.id, row["position"]) in existing_keys:
            row["errors"].append("This candidate already has an application for this position.")

        if not row["phone"]:
            row["errors"].append("Phone number is required.")
        if not row["position"]:
            row["errors"].append("Position is required.")
        if row["years_of_experience"] is None:
            row["errors"].append("Experience must be zero or a positive number.")
        if row["status"] not in valid_statuses:
            row["errors"].append("Status is not recognized.")
        if row["source"] is None:
            row["errors"].append("Source is not recognized.")
        row["source_label"] = source_labels.get(row["source"], "")

        row["is_valid"] = not row["errors"]

    return {
        "filename": uploaded_file.name,
        "sheet_name": worksheet.title,
        "rows": parsed_rows,
        "total_rows": len(parsed_rows),
        "valid_rows": sum(row["is_valid"] for row in parsed_rows),
        "invalid_rows": sum(not row["is_valid"] for row in parsed_rows),
    }


@transaction.atomic
def candidate_import_confirm(rows, *, user=None, assigned_recruiter=None):
    """Create one application per valid row, reusing the person if their email already exists."""
    valid_rows = [row for row in rows if row.get("is_valid")]
    emails = [row["email"] for row in valid_rows]
    company = get_user_company(user)
    candidates_by_email = _candidates_by_email(company, emails)
    existing_keys = _existing_application_keys(company, emails)
    jobs_by_title = _open_jobs_by_title(company)

    new_candidates = []
    applications = []
    skipped = 0
    for row in valid_rows:
        position = row["position"]
        job = jobs_by_title.get(position.lower())
        key = _application_key(row["email"], job and job.id, position)
        # Re-checked here because the data may have changed since the preview.
        if key in existing_keys:
            skipped += 1
            continue
        existing_keys.add(key)

        candidate = candidates_by_email.get(row["email"])
        if candidate is None:
            candidate = Candidate(
                company=company,
                full_name=row["full_name"],
                email=row["email"],
                phone=row["phone"],
                years_of_experience=row["years_of_experience"],
                source=row["source"],
            )
            candidates_by_email[row["email"]] = candidate
            new_candidates.append(candidate)

        applications.append(
            Application(
                candidate=candidate,
                job=job,
                imported_position="" if job else position,
                status=row["status"],
                assigned_recruiter=assigned_recruiter,
            )
        )

    Candidate.objects.bulk_create(new_candidates)
    Application.objects.bulk_create(applications)
    ApplicationStatusChange.objects.bulk_create(
        ApplicationStatusChange(
            application=application,
            from_status="",
            to_status=application.status,
            changed_by=user if user and user.is_authenticated else None,
        )
        for application in applications
    )

    return {
        "imported": len(applications),
        "new_candidates": len(new_candidates),
        "skipped": skipped,
    }
